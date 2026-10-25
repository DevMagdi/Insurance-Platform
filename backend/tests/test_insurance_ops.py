"""Phase 1 completions + Phase 2: unified CRM contracts + insurance operations.

Offline unit coverage for the pure domain logic and the API contract; the
PostgreSQL behaviour (RLS, SECURITY DEFINER importers/converters, gateway
payment recording) is exercised by the integration suite in CI.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from app import celery_app as celery_module
from app.celery_app import celery_app
from app.domain.crm_module_schemas import CrmLeadCreate, CrmStageCreate
from app.domain.enums import InstallmentPlan
from app.domain.insurance_ops_schemas import BrokeragePolicyCreate
from app.domain.permissions import PERMISSIONS
from app.main import app
from app.models import Base, BrokeragePolicy, CrmLead, Tenant
from app.services.crm_pipeline import DEFAULT_STAGES, slugify_client_name
from app.services.insurance_ops import (
    classify_insurance_type,
    commission_amount,
    installment_schedule,
    settlement_status,
)
from app.services.integration_providers import (
    WhatsApp360DialogProvider,
    WhatsAppCloudProvider,
    build_provider,
    build_whatsapp_provider,
)
from app.services.payments import verify_fawry, verify_paymob
from app.services.tabular import normalize_mobile, parse_tabular


@pytest.mark.parametrize(
    ("plan", "count"),
    [
        (InstallmentPlan.MONTHLY, 12),
        (InstallmentPlan.QUARTERLY, 4),
        (InstallmentPlan.SEMI_ANNUAL, 2),
        (InstallmentPlan.CASH, 1),
    ],
)
def test_installment_plans_split_the_full_premium(plan: InstallmentPlan, count: int) -> None:
    specs = installment_schedule(Decimal("100001.00"), plan, date(2026, 1, 31))
    assert len(specs) == count
    assert sum(spec.amount for spec in specs) == Decimal("100001.00")
    assert [spec.sequence_no for spec in specs] == list(range(1, count + 1))
    assert specs[0].due_date == date(2026, 1, 31)
    with pytest.raises(ValueError, match="positive"):
        installment_schedule(Decimal("0"), plan, date(2026, 1, 1))


def test_monthly_schedule_clamps_month_ends() -> None:
    specs = installment_schedule(Decimal("1200"), InstallmentPlan.MONTHLY, date(2026, 1, 31))
    assert specs[1].due_date == date(2026, 2, 28)
    assert specs[-1].due_date == date(2026, 12, 31)


def test_commission_defaults_to_five_percent_and_settles() -> None:
    assert commission_amount(Decimal("200000")) == Decimal("10000.00")
    assert commission_amount(Decimal("200000"), Decimal("0.075")) == Decimal("15000.00")
    assert settlement_status(Decimal("100"), Decimal("0")) == "pending"
    assert settlement_status(Decimal("100"), Decimal("40")) == "partial"
    assert settlement_status(Decimal("100"), Decimal("100")) == "settled"


def test_default_pipeline_matches_the_spec_stages() -> None:
    codes = [code for code, _, _ in DEFAULT_STAGES]
    assert codes == [
        "new",
        "contacted",
        "meeting",
        "needs_info",
        "quote_sent",
        "negotiation",
        "won",
        "lost",
        "converted",
    ]
    with pytest.raises(ValidationError):
        CrmStageCreate(code="x_stage", name_ar="س", name_en="X", is_won=True, is_lost=True)


def test_lead_and_policy_schemas_are_strict() -> None:
    lead = CrmLeadCreate(
        full_name="Ali", company_name="Acme", phone="01012345678", source="whatsapp"
    )
    assert lead.stage_code == "new" and lead.phone == "+201012345678"
    with pytest.raises(ValidationError):
        CrmLeadCreate(full_name="Ali", phone="123")
    with pytest.raises(ValidationError):
        BrokeragePolicyCreate(
            corporate_client_id="11111111-1111-4111-8111-111111111111",
            insurer_company_id="22222222-2222-4222-8222-222222222222",
            policy_no="POL-1",
            insurance_type="health",
            starts_on=date(2026, 6, 1),
            ends_on=date(2026, 1, 1),
            premium=Decimal("1000"),
        )


def test_classifier_mobile_and_slug_helpers() -> None:
    assert classify_insurance_type("طلب تأمين طبي للموظفين") == "health"
    assert classify_insurance_type("motor fleet renewal") == "general"
    assert classify_insurance_type("تأمين طبي وسيارات") == "both"
    assert normalize_mobile("01012345678") == "+201012345678"
    assert normalize_mobile("+20 101 234 5678") == "+201012345678"
    assert normalize_mobile("123") is None
    assert slugify_client_name("Acme Trading Co.", "x") == "acme-trading-co"
    assert slugify_client_name("شركة عربية", "abcdef123456").startswith("client-")


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["الاسم", "الموبايل", "البريد", "كود الموظف"])
    sheet.append(["Mona Ali", "01012345678", "mona@example.com", "E-1"])
    sheet.append(["", "", "", ""])  # blank rows are skipped
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_excel_and_csv_import_parsing_maps_arabic_headers() -> None:
    rows = parse_tabular("employees.xlsx", _xlsx_bytes())
    assert rows == [
        {
            "full_name": "Mona Ali",
            "mobile": "01012345678",
            "email": "mona@example.com",
            "employee_code": "E-1",
        }
    ]
    rows = parse_tabular("leads.csv", b"company_name,mobile\nAcme,01012345678\n")
    assert rows[0]["company_name"] == "Acme"
    with pytest.raises(ValueError, match="xlsx"):
        parse_tabular("leads.pdf", b"nope")


def test_payment_signatures_fail_closed() -> None:
    assert not verify_paymob(None, {}, "abc")
    assert not verify_paymob("secret", {"success": "true"}, "deadbeef")
    assert not verify_fawry(None, "M1", "REF", Decimal("10"), "abc")
    assert not verify_fawry("secret", "M1", "REF", Decimal("10"), "deadbeef")


def test_whatsapp_bsp_selection() -> None:
    assert isinstance(build_whatsapp_provider({}), WhatsAppCloudProvider)
    assert isinstance(build_whatsapp_provider({"bsp": "360dialog"}), WhatsApp360DialogProvider)
    assert isinstance(build_provider("whatsapp_360dialog", {}), WhatsApp360DialogProvider)
    bsp = WhatsApp360DialogProvider({})
    assert bsp._messages_url("1", "v21.0") == "https://waba-v2.360dialog.io/messages"
    assert bsp._auth_headers("key")["D360-API-KEY"] == "key"


def test_permissions_celery_and_openapi_contract() -> None:
    assert {
        "crm.view",
        "crm.manage",
        "crm.convert",
        "policies.view",
        "policies.manage",
        "issuing.view",
        "issuing.manage",
        "imports.manage",
        "payments.manage",
    } <= PERMISSIONS
    tasks = {entry["task"] for entry in celery_app.conf.beat_schedule.values()}
    assert tasks == {
        "app.celery_app.crm_task_reminders",
        "app.celery_app.mark_overdue_installments",
        "app.celery_app.policy_renewal_alerts",
    }
    paths = app.openapi()["paths"]
    for required in (
        "/api/crm/leads",
        "/api/crm/leads/{lead_id}/convert",
        "/api/crm/pipeline",
        "/api/staff/insurance-ops/policies",
        "/api/staff/insurance-ops/issuing",
        "/api/staff/insurance-ops/imports/employees",
        "/api/staff/insurance-ops/client-folder/{corporate_client_id}",
        "/api/staff/insurance-ops/caller-lookup",
        "/api/staff/insurance-ops/insurer-performance",
        "/api/payments/{provider}/webhook",
    ):
        assert required in paths, required
    # One CRM only: nothing may live under the old split prefix.
    assert not any(path.startswith("/api/staff/crm") for path in paths)
    lead_post = paths["/api/crm/leads"]["post"]
    assert lead_post["x-required-permission"] == "crm.manage"


def test_crm_models_register_from_the_package_import() -> None:
    # `from app.models import Base` alone must register every Phase 2 table:
    # the RLS integration catalog compares metadata to the migrated schema
    # without importing the FastAPI app first.
    assert CrmLead.__tablename__ == "crm_leads"
    assert Tenant.__tablename__ == "tenants"
    assert BrokeragePolicy.__tablename__ == "brokerage_policies"
    for table in (
        "tenants",
        "crm_pipeline_stages",
        "crm_leads",
        "crm_tasks",
        "crm_activities",
        "issuing_checklist_templates",
        "issuing_requests",
        "issuing_checklist_items",
        "brokerage_policies",
        "policy_installments",
        "commission_settlements",
        "payment_transactions",
        "import_jobs",
    ):
        assert table in Base.metadata.tables, table
    for gone in ("leads", "pipeline_stages", "lead_activities"):
        assert gone not in Base.metadata.tables, gone
    assert "tenant_id" in Base.metadata.tables["corporate_clients"].columns
    assert "is_owner" in Base.metadata.tables["corporate_users"].columns


def test_celery_tasks_run_their_sql_through_the_worker_scope(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    statements: list[str] = []
    monkeypatch.setattr(celery_module, "_run", lambda sql: statements.append(sql) or 2)
    assert celery_module.mark_overdue_installments.run() == {"overdue": 2}
    assert celery_module.policy_renewal_alerts.run() == {"alerts": 2}
    assert "policy_installments" in statements[0]
    assert "brokerage_policies" in statements[1]

    class Result:
        def scalar(self) -> int:
            return 3

    class Connection:
        def __enter__(self) -> Connection:
            return self

        def __exit__(self, *_: object) -> None:
            return None

        def execute(self, *_: object) -> Result:
            return Result()

    class Engine:
        def connect(self) -> Connection:
            return Connection()

        def dispose(self) -> None:
            return None

    monkeypatch.setattr(celery_module, "_engine", lambda: Engine())
    assert celery_module.crm_task_reminders.run() == {"reminders": 2, "due_leads": 3}
    assert "crm_tasks" in statements[-1]
