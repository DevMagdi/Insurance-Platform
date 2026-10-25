"""Tabular (Excel/CSV) parsing shared by every spreadsheet import.

Database-free and defensive: spreadsheet content is never trusted,
headers are canonicalised through Arabic/English aliases, Egyptian
mobiles normalise to ``+20`` form, and row counts are capped. Used by the
CRM lead import (`app.services.crm_import`) and the corporate
employee/client imports in insurance operations.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable

_MOBILE_LOCAL = re.compile(r"^01[0-9]{9}$")
_MOBILE_INT = re.compile(r"^\+?20(1[0-9]{9})$")


def normalize_mobile(raw: str) -> str | None:
    """Normalize an Egyptian mobile from a spreadsheet to +20 form, or None."""
    cleaned = re.sub(r"[\s\-()]", "", raw.strip())
    if cleaned.endswith(".0"):  # Excel renders numbers as floats.
        cleaned = cleaned[:-2]
    if _MOBILE_LOCAL.fullmatch(cleaned):
        return f"+20{cleaned[1:]}"
    match = _MOBILE_INT.fullmatch(cleaned)
    if match:
        return f"+20{match.group(1)}"
    return None


# Arabic spreadsheet headers seen in client files, mapped to canonical keys.
HEADER_ALIASES: dict[str, str] = {
    "الاسم": "full_name",
    "اسم الموظف": "full_name",
    "اسم الشركة": "company_name",
    "الشركة": "company_name",
    "الموبايل": "mobile",
    "الهاتف": "mobile",
    "رقم الموبايل": "mobile",
    "البريد": "email",
    "البريد الإلكتروني": "email",
    "الرقم القومي": "national_id",
    "كود الموظف": "employee_code",
    "جهة الاتصال": "contact_name",
    "شخص الاتصال": "contact_name",
    "القيمة المتوقعة": "expected_premium",
    "القسط المتوقع": "expected_premium",
    "المصدر": "source",
    "ملاحظات": "notes",
    "name": "full_name",
    "phone": "mobile",
}

_MAX_ROWS = 10_000


def canonical_header(raw: object) -> str:
    header = str(raw or "").strip().lower()
    return HEADER_ALIASES.get(header, re.sub(r"\s+", "_", header))


def rows_from_matrix(matrix: Iterable[Iterable[object]]) -> list[dict[str, str]]:
    iterator = iter(matrix)
    headers = [canonical_header(cell) for cell in next(iterator, [])]
    if not any(headers):
        raise ValueError("The spreadsheet has no header row")
    rows: list[dict[str, str]] = []
    for raw_row in iterator:
        values = ["" if cell is None else str(cell).strip() for cell in raw_row]
        if not any(values):
            continue
        rows.append({h: v for h, v in zip(headers, values, strict=False) if h})
        if len(rows) > _MAX_ROWS:
            raise ValueError(f"Imports are limited to {_MAX_ROWS} rows per file")
    return rows


def parse_tabular(filename: str, content: bytes) -> list[dict[str, str]]:
    """Parse an .xlsx or .csv upload into header-keyed string rows."""
    lowered = filename.lower()
    if lowered.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return rows_from_matrix(csv.reader(io.StringIO(text)))
    if lowered.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            return rows_from_matrix(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    raise ValueError("Only .xlsx and .csv files can be imported")
