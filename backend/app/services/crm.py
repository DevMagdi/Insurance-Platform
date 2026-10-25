"""Pure domain logic for Sales CRM, policy lifecycle and Excel imports.

Everything in this module is database-free so it is unit-testable offline:
installment schedules, the ~5% brokerage commission, default pipeline stages,
Tabular (Excel/CSV) parsing and mobile normalisation, and the deterministic
insurance-type classifier used when no LLM provider is configured.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from app.domain.enums import InstallmentPlan

CENT = Decimal("0.01")
DEFAULT_COMMISSION_RATE = Decimal("0.05")

# Spec Phase 2.1 pipeline, in order. "ربح / خسارة" are two terminal stages and
# "محوّل لعميل" follows a win; tenants may add/reorder stages at runtime.
DEFAULT_STAGES: tuple[tuple[str, str, str], ...] = (
    ("new", "جديد", "New"),
    ("contacted", "تم التواصل", "Contacted"),
    ("meeting", "في ميتنج", "Meeting"),
    ("needs_info", "يحتاج معلومات / عرض سعر", "Needs info / quotation"),
    ("quote_sent", "عرض مرسل", "Quote sent"),
    ("negotiation", "تفاوض", "Negotiation"),
    ("won", "ربح", "Won"),
    ("lost", "خسارة", "Lost"),
    ("converted", "محوّل لعميل", "Converted to client"),
)

# (count, months between installments) per premium plan.
_PLAN_SHAPE: dict[InstallmentPlan, tuple[int, int]] = {
    InstallmentPlan.MONTHLY: (12, 1),
    InstallmentPlan.QUARTERLY: (4, 3),
    InstallmentPlan.SEMI_ANNUAL: (2, 6),
    InstallmentPlan.CASH: (1, 0),
}


def _add_months(day: date, months: int) -> date:
    month = day.month - 1 + months
    year = day.year + month // 12
    month = month % 12 + 1
    # Clamp to the last day of the target month (Jan 31 + 1mo -> Feb 28/29).
    next_month = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)
    last_day = (next_month - date(year, month, 1)).days
    return date(year, month, min(day.day, last_day))


@dataclass(frozen=True)
class InstallmentSpec:
    sequence_no: int
    due_date: date
    amount: Decimal


def installment_schedule(
    premium: Decimal, plan: InstallmentPlan, starts_on: date
) -> list[InstallmentSpec]:
    """Split ``premium`` per plan; the final installment absorbs rounding."""
    if premium <= 0:
        raise ValueError("Premium must be positive")
    count, step = _PLAN_SHAPE[plan]
    each = (premium / count).quantize(CENT, rounding=ROUND_HALF_UP)
    specs: list[InstallmentSpec] = []
    allocated = Decimal("0")
    for index in range(count):
        amount = premium - allocated if index == count - 1 else each
        allocated += amount
        specs.append(InstallmentSpec(index + 1, _add_months(starts_on, step * index), amount))
    return specs


def commission_amount(premium: Decimal, rate: Decimal = DEFAULT_COMMISSION_RATE) -> Decimal:
    """Brokerage commission on a premium (spec default ≈5%)."""
    if premium < 0 or rate < 0 or rate > 1:
        raise ValueError("Premium and commission rate must be within range")
    return (premium * rate).quantize(CENT, rounding=ROUND_HALF_UP)


def settlement_status(expected: Decimal, received: Decimal) -> str:
    if received <= 0:
        return "pending"
    if received < expected:
        return "partial"
    return "settled"


_HEALTH_WORDS = ("طبي", "صحي", "مستشفى", "عيادة", "health", "medical", "clinic", "hospital")
_GENERAL_WORDS = ("سيارة", "سيارات", "حريق", "ممتلكات", "general", "motor", "fire", "property")


def classify_insurance_type(text: str) -> str:
    """Deterministic طبي/عام classifier (AI fallback when no LLM is wired).

    Mirrors the ticket auto-classification in the Phase 2 AI spec; returns
    ``both`` when medical and general signals appear together.
    """
    lowered = text.lower()
    health = any(word in lowered for word in _HEALTH_WORDS)
    general = any(word in lowered for word in _GENERAL_WORDS)
    if health and general:
        return "both"
    return "health" if health else "general"


def slugify_company(name: str, fallback: str) -> str:
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.lower()).strip("-")[:60]
    return slug or f"client-{fallback[:12]}"


_MOBILE_LOCAL = re.compile(r"^01[0-9]{9}$")
_MOBILE_INT = re.compile(r"^\+?20(1[0-9]{9})$")


def normalize_mobile(raw: str) -> str | None:
    """Normalize an Egyptian mobile from a spreadsheet to +20 form, or None."""
    cleaned = re.sub(r"[\s\-()]", "", raw.strip())
    if cleaned.endswith(".0"):  # Excel renders numbers as floats.
        cleaned = cleaned[:-2]
    if _MOBILE_LOCAL.fullmatch(cleaned):
        return f"+20{cleaned[1:]}"
    match = _MOBILE_INT.fullmatch(cleaned)
    if match:
        return f"+20{match.group(1)}"
    return None


# Arabic spreadsheet headers seen in client files, mapped to canonical keys.
HEADER_ALIASES: dict[str, str] = {
    "الاسم": "full_name",
    "اسم الموظف": "full_name",
    "اسم الشركة": "company_name",
    "الشركة": "company_name",
    "الموبايل": "mobile",
    "الهاتف": "mobile",
    "رقم الموبايل": "mobile",
    "البريد": "email",
    "البريد الإلكتروني": "email",
    "الرقم القومي": "national_id",
    "كود الموظف": "employee_code",
    "جهة الاتصال": "contact_name",
    "شخص الاتصال": "contact_name",
    "القيمة المتوقعة": "expected_premium",
    "القسط المتوقع": "expected_premium",
    "المصدر": "source",
    "ملاحظات": "notes",
    "name": "full_name",
    "phone": "mobile",
}

_MAX_ROWS = 10_000


def _canonical_header(raw: object) -> str:
    header = str(raw or "").strip().lower()
    return HEADER_ALIASES.get(header, re.sub(r"\s+", "_", header))


def _rows_from_matrix(matrix: Iterable[Iterable[object]]) -> list[dict[str, str]]:
    iterator = iter(matrix)
    headers = [_canonical_header(cell) for cell in next(iterator, [])]
    if not any(headers):
        raise ValueError("The spreadsheet has no header row")
    rows: list[dict[str, str]] = []
    for raw_row in iterator:
        values = ["" if cell is None else str(cell).strip() for cell in raw_row]
        if not any(values):
            continue
        rows.append({h: v for h, v in zip(headers, values, strict=False) if h})
        if len(rows) > _MAX_ROWS:
            raise ValueError(f"Imports are limited to {_MAX_ROWS} rows per file")
    return rows


def parse_tabular(filename: str, content: bytes) -> list[dict[str, str]]:
    """Parse an .xlsx or .csv upload into header-keyed string rows."""
    lowered = filename.lower()
    if lowered.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return _rows_from_matrix(csv.reader(io.StringIO(text)))
    if lowered.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            return _rows_from_matrix(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    raise ValueError("Only .xlsx and .csv files can be imported")
